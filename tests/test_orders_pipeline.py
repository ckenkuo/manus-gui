# -*- coding: utf-8 -*-
"""订单登记管线【离线部分】单测：解析官方导出 xlsx、站点归一化、字段/判重列映射。

页面部分（翻页勾选、触发导出、抓图）不在此覆盖——那要真浏览器，靠实机 dry-run 验证。
这里钉死的是「导出文件 → 登记表行」这段纯数据变换，它决定会往 102MB 的表里写什么。

样本文件 workspace/orders_sample/sample_export.xlsx 是 gitignored 的实机导出（22 行 /
20 个订单），缺失时相关用例 skip，不让 CI 因为没样本就红。
"""
import asyncio
from pathlib import Path

import pytest

from app.orders import pipeline as P
from app.orders.pipeline import OrderRow

SAMPLE = Path("workspace/orders_sample/sample_export.xlsx")
needs_sample = pytest.mark.skipif(not SAMPLE.exists(), reason="缺实机导出样本")

# StoreA全球1 的真实表头（实测），列序刻意不连续：订单号在 C 而非 A
STORE_A_HEADER = {
    "A": "订单店铺", "B": "站点区分", "C": "订单号", "D": "尺码",
    "E": "平台物流跟踪号", "F": "国内发出时间", "G": "产品图片", "H": "产品图2",
    "I": "平台创建时间", "J": "采购日期", "K": "采购费用", "L": "平台成交价",
    "M": "Y2头程费用", "N": "采购订单号", "O": "物流情况",
}


def _order(**kw) -> OrderRow:
    base = dict(
        order_no="PO-045-119", site="哥伦比亚", sub_order_no="045-118",
        attrs="杏色 / 3-4Y", created_at="2026-07-27 10:16:38",
    )
    base.update(kw)
    return OrderRow(**base)


# ---- 站点归一化 / 空值 ------------------------------------------------------


def test_normalize_site_strips_trailing_zhan():
    """导出的「哥伦比亚站」要归一成表内主流写法「哥伦比亚」。"""
    assert P.normalize_site("哥伦比亚站") == "哥伦比亚"
    assert P.normalize_site("美国") == "美国"
    assert P.normalize_site(" 秘鲁站 ") == "秘鲁"
    # 单字「站」不该被吃成空串
    assert P.normalize_site("站") == "站"
    # `--` 是导出文件的空占位
    assert P.normalize_site("--") == ""
    assert P.normalize_site(None) == ""


def test_clean_normalizes_placeholder_empties():
    assert P._clean("--") == ""
    assert P._clean("—") == ""
    assert P._clean(" 杏色 / 3-4Y ") == "杏色 / 3-4Y"
    assert P._clean(0) == "0", "数字 0 是有效值，不能当空"


def test_to_jpeg_url():
    """Excel/WPS 不认 avif，嵌进去是空白格。"""
    assert P.to_jpeg_url(
        "https://img.kwcdn.com/a-goods.jpeg?imageView2/2/w/800/q/70/format/avif"
    ).endswith("format/jpeg")
    # 没带 format 参数的 imageView2 URL 要补上
    assert P.to_jpeg_url("https://img.kwcdn.com/a-goods.jpeg?imageView2/2/w/800").endswith(
        "/format/jpeg"
    )
    # 非 imageView2 的 URL 原样不动，别乱拼参数
    plain = "https://img.kwcdn.com/a-goods.jpeg"
    assert P.to_jpeg_url(plain) == plain
    assert P.to_jpeg_url("") == ""


# ---- 解析导出 xlsx ---------------------------------------------------------


@needs_sample
def test_parse_export_sample():
    """实机样本：22 行子订单级记录，字段对得上、`--` 已归一化为空。"""
    rows = P.parse_export_xlsx(str(SAMPLE))

    assert len(rows) == 22
    # 按子订单展开：20 个订单 → 22 行，说明有订单含多个子订单
    assert len({r.order_no for r in rows}) == 20

    # 不钉死具体单号：真实单号属业务数据，不入库。改为按不变量断言——
    # 取任一哥伦比亚单，校验字段形态而非字面值。
    o = next(r for r in rows if r.site == "哥伦比亚")
    assert not o.site.endswith("站"), "站点必须已去掉尾字「站」"
    assert o.status == "待发货"
    assert o.order_no.startswith("PO-") and o.sub_order_no
    assert o.attrs, "商品属性非空（形如「颜色 / 尺码」，也有仅款式无尺码的）"
    assert o.created_at.startswith("2026-") and len(o.created_at) == 19
    assert o.spu_id.isdigit()
    # 待发货态这些列在导出里全是 `--`
    assert (o.tracking_no, o.carrier, o.warehouse, o.shipped_at) == ("", "", "", "")
    # DOM 才有的字段，解析阶段还是空
    assert (o.image_url, o.deal_price, o.image_path) == ("", "", "")

    assert all(r.order_no and r.sub_order_no for r in rows), "订单号/子订单号不得为空"


def test_parse_export_rejects_missing_columns(tmp_path):
    """表头缺关键列直接抛错——字段对不上就别往登记表写。"""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["订单号", "站点", "订单状态"])  # 缺 子订单号/商品属性/订单创建时间
    ws.append(["PO-1", "美国站", "待发货"])
    p = tmp_path / "bad.xlsx"
    wb.save(p)

    with pytest.raises(ValueError, match="缺列"):
        P.parse_export_xlsx(str(p))


def test_parse_export_is_column_order_agnostic(tmp_path):
    """按表头标题取列，不认列序：导出字段设置一改列序就会变。"""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["商品属性", "订单创建时间", "子订单号", "站点", "订单号", "运单号"])
    ws.append(["粉色 / 5-6Y", "2026-07-27 09:25:52", "045-999", "秘鲁站", "PO-159-1", "--"])
    ws.append([None, None, None, None, None, None])  # 全空行要被跳过
    p = tmp_path / "reordered.xlsx"
    wb.save(p)

    rows = P.parse_export_xlsx(str(p))
    assert len(rows) == 1
    assert (rows[0].order_no, rows[0].site, rows[0].attrs) == ("PO-159-1", "秘鲁", "粉色 / 5-6Y")
    assert rows[0].tracking_no == ""


def test_summarize_purchases_groups_same_spu_and_sku():
    orders = [
        _order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A", sku_id="SKU-A",
               sku_code="CODE-A", qty="1", goods_name="连衣裙", attrs="杏色 / 3-4Y"),
        _order(order_no="PO-2", sub_order_no="045-2", spu_id="SPU-A", sku_id="SKU-A",
               sku_code="CODE-A", qty="2", goods_name="连衣裙", attrs="杏色 / 3-4Y"),
        _order(order_no="PO-3", sub_order_no="045-3", spu_id="SPU-A", sku_id="SKU-B",
               sku_code="CODE-B", qty="1", goods_name="连衣裙", attrs="粉色 / 5-6Y"),
    ]

    groups = P.summarize_purchases(orders)

    assert len(groups) == 2
    repeated = next(group for group in groups if group["sku_id"] == "SKU-A")
    assert repeated["total_qty"] == 3
    assert repeated["order_count"] == 2 and repeated["sub_order_count"] == 2
    assert repeated["is_repeated"] is True
    assert repeated["order_nos"] == ["PO-1", "PO-2"]


def test_export_purchase_summary_contains_summary_and_details(tmp_path):
    orders = [
        _order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A", sku_id="SKU-A", qty="1"),
        _order(order_no="PO-2", sub_order_no="045-2", spu_id="SPU-A", sku_id="SKU-A", qty="2"),
    ]

    result = P.export_purchase_summary(orders, str(tmp_path), "20260729")

    assert result["rows"] == 2 and result["groups"] == 1
    assert result["repeated_groups"] == 1 and result["total_qty"] == 3
    import openpyxl

    workbook = openpyxl.load_workbook(result["file"], read_only=True, data_only=True)
    assert workbook.sheetnames == ["商品汇总", "SPU_SKU汇总", "订单明细"]
    product_rows = list(workbook["商品汇总"].iter_rows(values_only=True))
    summary_rows = list(workbook["SPU_SKU汇总"].iter_rows(values_only=True))
    detail_rows = list(workbook["订单明细"].iter_rows(values_only=True))
    workbook.close()
    # 商品级放第一张：采购按商品链接下单，先看这款要买哪些码各几件
    # 前两列固定是「主图」+ 人工勾的「是否采购完成」（2026-07-30 口径）
    assert product_rows[0][0:8] == (
        "主图", "是否采购完成", "需多规格", "商品名称", "SPU ID", "规格数",
        "采购总件数", "采购清单",
    )
    assert summary_rows[0][0:5] == (
        "主图", "是否采购完成", "是否重复采购", "SPU ID", "SKU ID",
    )
    assert summary_rows[1][0:5] == (None, None, "是", "SPU-A", "SKU-A")
    # 时间列只保留一个「创建时间」＝该组最早下单时间，不再出最早/最晚两列
    assert "创建时间" in product_rows[0] and "最晚下单时间" not in product_rows[0]
    assert len(detail_rows) == 3


def test_export_purchase_summary_uses_earliest_created_time(tmp_path):
    """合并成的那一个「创建时间」取该组【最早】下单时间：采购紧急度看它。"""
    orders = [
        _order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A", sku_id="SKU-A",
               qty="1", created_at="2026-07-28 09:00:00"),
        _order(order_no="PO-2", sub_order_no="045-2", spu_id="SPU-A", sku_id="SKU-A",
               qty="1", created_at="2026-07-26 08:00:00"),
    ]

    result = P.export_purchase_summary(orders, str(tmp_path), "20260730")

    import openpyxl

    workbook = openpyxl.load_workbook(result["file"], read_only=True, data_only=True)
    rows = list(workbook["商品汇总"].iter_rows(values_only=True))
    workbook.close()
    assert rows[1][rows[0].index("创建时间")] == "2026-07-26 08:00:00"


def test_export_purchase_summary_embeds_images_and_fits_one_screen(tmp_path):
    """两张汇总表逐行嵌主图；列宽合计不超过一屏可用宽度（只上下滚、不左右滚）。"""
    from PIL import Image

    img = tmp_path / "045-1.jpg"
    Image.new("RGB", (800, 800), (200, 120, 80)).save(img)
    orders = [
        _order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A", sku_id="SKU-A",
               qty="2", image_path=str(img)),
        # 没下到图的那条照常出行，只是图那格空着
        _order(order_no="PO-2", sub_order_no="045-2", spu_id="SPU-B", sku_id="SKU-B",
               qty="1", image_path=""),
    ]

    result = P.export_purchase_summary(orders, str(tmp_path), "20260730")

    assert result["product_images"] == 1 and result["sku_images"] == 1
    import openpyxl

    workbook = openpyxl.load_workbook(result["file"])
    avail = P.screen_client_px()
    for name, height in (("商品汇总", 60.0), ("SPU_SKU汇总", 60.0), ("订单明细", 25.0)):
        ws = workbook[name]
        total = sum(
            ws.column_dimensions[ws.cell(1, i).column_letter].width * 7 + 5
            for i in range(1, ws.max_column + 1)
        )
        assert total <= avail + 1, f"{name} 列宽合计 {total} 超出一屏 {avail}"
        assert ws.row_dimensions[2].height == height
        assert ws["C2"].alignment.wrap_text is True
        assert ws["C2"].alignment.vertical == "center"
        assert ws.freeze_panes == "A2"
    assert len(workbook["商品汇总"]._images) == 1
    assert len(workbook["SPU_SKU汇总"]._images) == 1
    assert len(workbook["订单明细"]._images) == 0, "明细是分单用的长表，不放图"
    # 锚定必须是「随单元格移动并调整大小」：这两张表开着筛选，浮动图筛完会糊在剩下的行上
    anchor = workbook["SPU_SKU汇总"]._images[0].anchor
    assert anchor.editAs == "twoCell"
    assert (anchor._from.col, anchor._from.row) == (0, 1), "落在 A2"
    assert (anchor.to.col, anchor.to.row) == (0, 1), "两个锚点同格，不跨到右边列"
    # 「是否采购完成」给是/否下拉，防手输五花八门的写法
    dv = workbook["商品汇总"].data_validations.dataValidation
    assert [(str(v.sqref), v.formula1) for v in dv] == [("B2:B3", '"是,否"')]
    workbook.close()


def test_export_purchase_summary_tolerates_missing_image_file(tmp_path):
    """图文件被删/路径失效只该那一格空着，不能让整份统计导不出来。"""
    orders = [_order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A",
                     sku_id="SKU-A", qty="1", image_path=str(tmp_path / "没有这个.jpg"))]

    result = P.export_purchase_summary(orders, str(tmp_path), "20260730")

    assert result["product_images"] == 0 and Path(result["file"]).exists()


def test_screen_client_px_reserves_room_and_has_floor():
    """列宽预算＝屏宽减去行号列/滚动条/边框的余量，且有下限兜住异常小的屏。"""
    assert P.screen_client_px(reserve=0) - P.screen_client_px(reserve=130) == 130
    assert P.screen_client_px(reserve=99999) == 800


def test_summarize_products_merges_sizes_of_same_goods():
    """同商品不同尺码必须并成一组、尺码收进 variants——采购是按商品链接下单的。

    这是实机撞出来的：一条公主裙 5 个尺码，按 (SPU,SKU) 聚合成了 5 组，看的人得自己
    在表里认哪几行是同一款，正好是这份统计该替他做的事。
    """
    orders = [
        _order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A", sku_id="SKU-3-4",
               qty="2", goods_name="公主裙", attrs="杏色 / 3-4Y"),
        _order(order_no="PO-2", sub_order_no="045-2", spu_id="SPU-A", sku_id="SKU-5-6",
               qty="1", goods_name="公主裙", attrs="杏色 / 5-6Y"),
        _order(order_no="PO-3", sub_order_no="045-3", spu_id="SPU-B", sku_id="SKU-B",
               qty="1", goods_name="牛仔裤", attrs="蓝色 / 5-6Y"),
    ]

    products = P.summarize_products(orders)

    assert len(products) == 2, "同 SPU 不同尺码要并成一个商品"
    dress = next(p for p in products if p["spu_id"] == "SPU-A")
    assert dress["variant_count"] == 2 and dress["total_qty"] == 3
    assert dress["is_multi"] is True
    assert [v["attrs"] for v in dress["variants"]] == ["杏色 / 3-4Y", "杏色 / 5-6Y"]
    assert [v["qty"] for v in dress["variants"]] == [2, 1]
    assert dress["order_nos"] == ["PO-1", "PO-2"]
    jeans = next(p for p in products if p["spu_id"] == "SPU-B")
    assert jeans["is_multi"] is False and jeans["variant_count"] == 1


def test_summarize_products_falls_back_to_goods_name_without_spu():
    """SPU 缺失时按商品名并，不要每条各成一组；不同商品仍要分开。"""
    orders = [
        _order(order_no="PO-1", sub_order_no="045-1", spu_id="", sku_id="S1",
               qty="1", goods_name="无SPU裙", attrs="3-4Y"),
        _order(order_no="PO-2", sub_order_no="045-2", spu_id="", sku_id="S2",
               qty="1", goods_name="无SPU裙", attrs="5-6Y"),
        _order(order_no="PO-3", sub_order_no="045-3", spu_id="", sku_id="S3",
               qty="1", goods_name="另一款", attrs="7-8Y"),
    ]

    products = P.summarize_products(orders)

    assert len(products) == 2
    assert next(p for p in products if "无SPU裙" in p["goods_names"])["variant_count"] == 2


def test_summarize_products_merges_same_size_across_orders():
    """同商品同尺码被多张单买到 → 件数累加成一个规格，不重复列。"""
    orders = [
        _order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A", sku_id="SKU-A",
               qty="1", goods_name="连衣裙", attrs="杏色 / 3-4Y"),
        _order(order_no="PO-2", sub_order_no="045-2", spu_id="SPU-A", sku_id="SKU-A",
               qty="2", goods_name="连衣裙", attrs="杏色 / 3-4Y"),
    ]

    products = P.summarize_products(orders)

    assert len(products) == 1
    assert products[0]["variant_count"] == 1
    assert products[0]["variants"][0]["qty"] == 3
    assert products[0]["variants"][0]["order_nos"] == ["PO-1", "PO-2"]
    assert products[0]["is_multi"] is True, "同规格多单也要拎出来（要合单采购）"


def test_export_purchase_markdown_groups_sizes_under_one_product(tmp_path):
    """md 要按商品分小节，同款的所有尺码列在一张表里——打开一次链接买齐。"""
    orders = [
        _order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A", sku_id="SKU-3-4",
               qty="2", goods_name="公主裙", attrs="杏色 / 3-4Y"),
        _order(order_no="PO-2", sub_order_no="045-2", spu_id="SPU-A", sku_id="SKU-5-6",
               qty="1", goods_name="公主裙", attrs="杏色 / 5-6Y"),
        _order(order_no="PO-3", sub_order_no="045-3", spu_id="SPU-B", sku_id="SKU-B",
               qty="1", goods_name="牛仔裤", attrs="蓝色 / 5-6Y"),
    ]

    result = P.export_purchase_markdown(orders, str(tmp_path), "20260729_101530")

    assert result["rows"] == 3 and result["products"] == 2
    assert result["multi_products"] == 1 and result["total_qty"] == 4
    assert result["variants"] == 3
    path = Path(result["md_file"])
    assert path.name == "新增订单采购统计_20260729_101530.md"
    text = path.read_text(encoding="utf-8")
    assert "本批待登记子订单：3 条" in text
    assert "涉及商品（SPU）：2 个" in text
    assert "一、需一次买多个码数" in text and "二、单规格单次采购" in text
    assert "三、子订单明细" in text
    # 公主裙自成一节，两个尺码同表；牛仔裤只有一个规格，进第二段
    multi = text.split("一、需一次买多个码数")[1].split("## 二、")[0]
    assert "### 公主裙" in multi and "SPU `SPU-A`" in multi
    assert "2 个规格" in multi
    assert "杏色 / 3-4Y" in multi and "杏色 / 5-6Y" in multi
    assert "牛仔裤" not in multi, "单规格商品不进第一段"
    detail = text.split("三、子订单明细")[1]
    assert "PO-1" in detail and "PO-2" in detail and "PO-3" not in detail


def test_export_purchase_markdown_escapes_pipe(tmp_path):
    """商品名带竖线会截断表格列，必须转义。"""
    orders = [_order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A",
                     sku_id="SKU-A", qty="1", goods_name="连衣裙|夏款")]

    result = P.export_purchase_markdown(orders, str(tmp_path), "20260729")

    assert "连衣裙\\|夏款" in Path(result["md_file"]).read_text(encoding="utf-8")


# ---- 文件名带店铺名：多店铺采购时要能一眼分辨 -------------------------------


def test_purchase_file_stem_prefixes_store():
    """店铺名放最前面，同店的文件按名称排序自然聚在一起。"""
    assert P.purchase_file_stem("新增订单采购汇总", "StoreA", "20260805_101530") == \
        "StoreA_新增订单采购汇总_20260805_101530"


def test_purchase_file_stem_strips_illegal_chars():
    """店名是页面自由文本，带 Windows 非法字符会让写文件直接失败。"""
    assert P.purchase_file_stem("汇总", 'A/B:C*?"<>|D', "20260805") == "ABCD_汇总_20260805"


def test_purchase_file_stem_without_store_keeps_old_format():
    """没识别到店铺时退回原格式，不留「_」空占位。"""
    assert P.purchase_file_stem("汇总", "", "20260805") == "汇总_20260805"
    assert P.purchase_file_stem("汇总", "  ", "20260805") == "汇总_20260805"


def test_export_purchase_files_carry_store_name(tmp_path):
    """xlsx 与 md 的文件名都带店铺；md 正文抬头也写店铺（会被整段贴去对量）。"""
    orders = [_order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A",
                     sku_id="SKU-A", qty="1")]

    xlsx = P.export_purchase_summary(orders, str(tmp_path), "20260805", "StoreB")
    md = P.export_purchase_markdown(orders, str(tmp_path), "20260805", "StoreB")

    assert Path(xlsx["file"]).name == "StoreB_新增订单采购汇总_20260805.xlsx"
    assert Path(md["md_file"]).name == "StoreB_新增订单采购统计_20260805.md"
    assert "- 店铺：StoreB" in Path(md["md_file"]).read_text(encoding="utf-8")


def test_export_purchase_markdown_omits_store_line_when_unknown(tmp_path):
    """店铺为空时不留「- 店铺：」空行。"""
    orders = [_order(order_no="PO-1", sub_order_no="045-1", spu_id="SPU-A",
                     sku_id="SKU-A", qty="1")]

    md = P.export_purchase_markdown(orders, str(tmp_path), "20260805")

    assert "店铺：" not in Path(md["md_file"]).read_text(encoding="utf-8")


# ---- 增量早停判据（纯函数，不需要浏览器）-----------------------------------


def _page(*items) -> list:
    """构造一页的 [{order_no, created_at}]，items 传 (订单号, 时间) 或只传订单号。"""
    out = []
    for it in items:
        no, t = it if isinstance(it, tuple) else (it, "")
        out.append({"order_no": no, "created_at": t})
    return out


def test_stop_when_whole_page_known_after_seeing_new():
    """整页全已登记 + 本批见过新单 → 早停（正常追上的情形）。"""
    stop, reason = P.should_stop_incremental(
        _page("PO-1", "PO-2"), {"PO-1", "PO-2"}, seen_new=True
    )
    assert stop is True and "追上" in reason


def test_no_stop_while_page_is_all_new():
    """本页一条已登记的都没有 → 还没追上，继续翻。"""
    stop, reason = P.should_stop_incremental(
        _page("PO-new1", "PO-new2"), {"PO-1"}, seen_new=False
    )
    assert stop is False and reason == ""


def test_stop_on_first_page_all_known_means_no_new_orders():
    """第一页就全已登记（还没见过新单）→ 停，判定为「本批无新单」。"""
    stop, reason = P.should_stop_incremental(
        _page("PO-1", "PO-2"), {"PO-1", "PO-2"}, seen_new=False
    )
    assert stop is True and "未发现新单" in reason


def test_stop_on_single_known_order_amid_new_ones():
    """稀疏水位：整页只有 1 条已登记、前后都是新单 → 照样停。

    2026-08-04 实机回归：新建 Sheet 只登记过 1 个订单号，它出现在第 4 页中间，
    前面的更新、后面的更旧，都没登记过。旧判据要求「整页全已登记」，水位只有 1 条时
    恒不成立，还会被交错检测判成「排序可能被改」而退全量，实际翻了 63 页。
    """
    stop, reason = P.should_stop_incremental(
        _page("PO-new1", "PO-new2", "PO-1", "PO-new3", "PO-new4"),
        {"PO-1"}, seen_new=True,
    )
    assert stop is True and "追上" in reason and "PO-1" in reason


def test_stop_when_known_order_is_last_on_page():
    """已登记单出现在页尾 → 停（它之后的都更旧，不是新单）。"""
    stop, reason = P.should_stop_incremental(
        _page("PO-new1", "PO-new2", "PO-1"), {"PO-1"}, seen_new=True
    )
    assert stop is True and "追上" in reason


def test_empty_page_does_not_stop():
    """读不到订单号时绝不早停——宁可多翻几页也不能因为读不到就停。"""
    stop, reason = P.should_stop_incremental([], {"PO-1"}, seen_new=True)
    assert stop is False and "读不到" in reason


def test_empty_watermark_never_stops():
    """水位为空（首次登记）→ 全是新单，自然翻到底。"""
    stop, _ = P.should_stop_incremental(_page("PO-1", "PO-2"), set(), seen_new=False)
    assert stop is False


def test_check_list_sort_url_accepts_desc():
    """sortType=1（创建时间新→旧）→ 通过。"""
    url = "https://x/orders.html?fulfillmentMode=0&queryType=2&sortType=1&timeZone=UTC%2B8"
    assert P.check_list_sort_url(url) == ""


def test_check_list_sort_url_rejects_wrong_sort():
    """sortType 不是 1 → 返回中止文案，且要指明怎么改、怎么绕。

    这是静默漏采的唯一确定性护栏：排序反了会停在第 1 页，汇总却显示「本批无新单」。
    """
    err = P.check_list_sort_url("https://x/orders.html?queryType=2&sortType=2")

    assert "sortType=2" in err
    assert "config.toml" in err, "要告诉操作者去哪儿改"
    assert "--no-incremental" in err, "要给出不改配置也能跑的出路"


def test_check_list_sort_url_rejects_missing_sort():
    """URL 里压根没有 sortType → 同样中止：无法确认排序就不能信早停。"""
    err = P.check_list_sort_url("https://x/orders.html?queryType=2")

    assert "没有 sortType" in err and "config.toml" in err


def test_check_list_sort_url_ignores_empty_url():
    """list_url 为空 → 交给调用方的必填校验报，这里不重复报。"""
    assert P.check_list_sort_url("") == ""


def test_check_created_desc_accepts_descending():
    """页内倒序且不比上一页更新 → 通过。"""
    page = _page(("PO-2", "2026-07-28 10:00:00"), ("PO-1", "2026-07-27 09:00:00"))
    assert P.check_created_desc("2026-07-29 12:00:00", page) == ""


def test_check_created_desc_flags_ascending_within_page():
    """页内出现升序 → 报排序异常。"""
    page = _page(("PO-1", "2026-07-27 09:00:00"), ("PO-2", "2026-07-28 10:00:00"))
    assert "倒序" in P.check_created_desc("", page)


def test_check_created_desc_flags_newer_than_prev_page():
    """本页出现比上一页更新的时间 → 报排序异常。"""
    page = _page(("PO-9", "2026-07-30 10:00:00"))
    assert "更新的创建时间" in P.check_created_desc("2026-07-28 10:00:00", page)


def test_check_created_desc_degrades_without_times():
    """列表页读不到创建时间 → 返回空串，静默降级到无排序护栏。"""
    assert P.check_created_desc("2026-07-28 10:00:00", _page("PO-1", "PO-2")) == ""


# ---- 列映射 / 判重键 -------------------------------------------------------


def test_build_row_values_uses_real_header():
    """按真实表头生成 {列字母: 值}；人工填的列一律不碰。"""
    o = _order(tracking_no="", deal_price="63.99")
    values = P.build_row_values(o, STORE_A_HEADER, store_value="StoreA全球")

    assert values["A"] == "StoreA全球"      # 配置常量，≠ 登录店铺名
    assert values["B"] == "哥伦比亚"
    assert values["C"] == "PO-045-119"
    assert values["D"] == "杏色 / 3-4Y"     # 商品属性 → 尺码
    assert values["E"] == ""                # 待发货无运单号
    assert values["I"] == "2026-07-27 10:16:38"
    assert values["L"] == "63.99"           # 成交价 best-effort
    # 人工后续填的列：管线不得写入，连空串都不占
    for col in ("F", "J", "K", "M", "N", "O"):
        assert col not in values, f"{col}({STORE_A_HEADER[col]}) 应由人工填，管线不该碰"
    # 图片列不走 values（走 DISPIMG 嵌图）
    assert "G" not in values and "H" not in values


def test_build_row_values_writes_quantity_as_number():
    """有「数量」列就写应履约件数，且必须是数字——这一列要能求和、筛多件单。"""
    header = dict(STORE_A_HEADER)
    header["P"] = "数量"
    values = P.build_row_values(_order(qty="2"), header, store_value="StoreA全球")

    assert values["P"] == 2 and isinstance(values["P"], int)


def test_build_row_values_quantity_column_aliases():
    """用户可能早先手工加过「件数」这类写法，一并认下来，免得插出第二列同义列。"""
    for title in ("件数", "商品数量", "采购件数", "应履约件数"):
        header = dict(STORE_A_HEADER, P=title)
        assert P.build_row_values(_order(qty="3"), header, "S")["P"] == 3


def test_quantity_number_keeps_dirty_text_and_skips_empty():
    """空值不落单元格；解析不出的脏值留原文给人看，不要静默写 0。"""
    assert P._qty_number("") == "" and P._qty_number("--") == ""
    assert P._qty_number("2") == 2
    assert P._qty_number("1.5") == 1.5
    assert P._qty_number("两件") == "两件"


def test_dedupe_key_unaffected_by_quantity_column():
    """判重键是订单号+尺码，加了数量列不能改变键——否则历史行会被判成新行重写。"""
    header = dict(STORE_A_HEADER, P="数量")
    o = _order(qty="2")

    assert P.dedupe_key(o, header, ["订单号", "尺码"], "StoreA全球") == P.dedupe_key(
        o, STORE_A_HEADER, ["订单号", "尺码"], "StoreA全球"
    )


def test_image_column_prefers_exact_title():
    """`StoreA全球1` 同时有「产品图片」和「产品图2」，只能写前者。"""
    assert P.image_column(STORE_A_HEADER) == "G"
    # 只有「产品图2」这类变体时退回首个含「产品图」的列
    assert P.image_column({"C": "订单号", "E": "产品图2"}) == "E"
    assert P.image_column({"A": "订单号"}) is None


def test_resolve_title_column_handles_alias_and_order():
    """字段有多种标题写法（` StoreD` 表用「站点」而非「站点区分」）。"""
    assert P.resolve_title_column(STORE_A_HEADER, {"站点区分", "站点"}) == "B"
    assert P.resolve_title_column({"B": "站点", "C": "订单号"}, {"站点区分", "站点"}) == "B"
    assert P.resolve_title_column(STORE_A_HEADER, "不存在的列") is None
    # 多列命中时取列序最靠前的那个（AA 排在 B 之后，不能按字符串比较）
    assert P.resolve_title_column({"AA": "订单号", "B": "订单号"}, "订单号") == "B"


def test_dedupe_key_matches_workbook_reading():
    """判重键 = (订单号, 尺码)，与 existing_key_tuples 读表口径一致（strip 后比较）。"""
    o = _order(attrs=" 杏色 / 3-4Y ")
    key = P.dedupe_key(o, STORE_A_HEADER, ["订单号", "尺码"], "StoreA全球")
    assert key == ("PO-045-119", "杏色 / 3-4Y")

    # 同订单不同尺码必须是两个键，否则同订单的其余子订单会被当重复丢掉
    k2 = P.dedupe_key(_order(attrs="杏色 / 9-12M"), STORE_A_HEADER, ["订单号", "尺码"], "StoreA全球")
    assert k2 != key


def test_dedupe_key_returns_none_when_column_absent():
    """判重列在该 Sheet 找不到 → None，调用方据此跳过整个 Sheet，而不是硬写重复行。"""
    header = {"A": "订单店铺", "B": "订单号"}  # 没有「尺码」列
    assert P.dedupe_key(_order(), header, ["订单号", "尺码"], "StoreA") is None


# ---- Sheet 映射 ------------------------------------------------------------


SHEET_MAP = [
    {"store": "StoreA", "sites": ["哥伦比亚", "秘鲁"], "sheet": "StoreA全球1",
     "store_value": "StoreA全球"},
    {"store": "StoreC", "sites": ["美国"], "sheet": "StoreC美国",
     "store_value": "StoreC"},
]


def test_resolve_sheet_matches_store_and_site():
    got = P.resolve_sheet("StoreA", "哥伦比亚站", SHEET_MAP)
    assert got and got["sheet"] == "StoreA全球1" and got["store_value"] == "StoreA全球"
    # 店铺名带后缀也要认（页面右上角可能显示 "StoreA Store"）
    assert P.resolve_sheet("StoreA Store", "秘鲁", SHEET_MAP)["sheet"] == "StoreA全球1"
    assert P.resolve_sheet("StoreC", "美国", SHEET_MAP)["sheet"] == "StoreC美国"


def test_resolve_sheet_returns_none_when_unmapped():
    """未命中映射必须返回 None（调用方跳过并报告），绝不臆测落点。

    StoreA + 美国 就是典型陷阱：`StoreA牛仔裤` 虽是 StoreA 的美国单，但列结构完全不同
    （有 Y1/Y2、XM、货号），不能想当然当作 StoreA 美国站的落点。
    """
    assert P.resolve_sheet("StoreA", "美国", SHEET_MAP) is None
    assert P.resolve_sheet("陌生店铺", "哥伦比亚", SHEET_MAP) is None
    assert P.resolve_sheet("StoreA", "哥伦比亚", []) is None


# ---- join 图片 / 成交价 ----------------------------------------------------


def test_join_images_by_sub_order():
    """图与成交价按子订单号 join；缺图只统计不报错（图是辅助字段）。"""
    a = _order(sub_order_no="045-1")
    b = _order(sub_order_no="045-2")
    c = _order(sub_order_no="045-3")
    images = {
        "045-1": {"url": "https://img.kwcdn.com/1.jpeg", "price": "63.99", "loaded": True},
        "045-2": {"url": "https://img.kwcdn.com/2.jpeg", "price": "", "loaded": True},
    }

    stat = P.join_images([a, b, c], images)

    assert (a.image_url, a.deal_price) == ("https://img.kwcdn.com/1.jpeg", "63.99")
    assert (b.image_url, b.deal_price) == ("https://img.kwcdn.com/2.jpeg", "")
    assert (c.image_url, c.deal_price) == ("", ""), "没抓到图的行保持空，不得串到别人的图"
    assert stat == {"image": 2, "price": 1, "miss": 1}


def test_download_images_tolerates_single_failure(tmp_path, monkeypatch):
    """单张图下载失败只记 warning，该行不带图入库，不拖垮整批。"""
    calls = []

    def fake_dl(url, dst, *a, **kw):
        calls.append(url)
        if "bad" in url:
            raise RuntimeError("connection reset")
        Path(dst).write_bytes(b"\xff\xd8jpg")

    monkeypatch.setattr(P, "_download_main_image", fake_dl)
    ok_row = _order(sub_order_no="045-1", image_url="https://x/good.jpeg")
    bad_row = _order(sub_order_no="045-2", image_url="https://x/bad.jpeg")
    no_img = _order(sub_order_no="045-3")

    stat = P.download_images([ok_row, bad_row, no_img], str(tmp_path / "img"))

    assert stat == {"ok": 1, "fail": 1}
    assert ok_row.image_path and Path(ok_row.image_path).exists()
    assert bad_row.image_path == "" and no_img.image_path == ""
    assert len(calls) == 2, "没有 image_url 的行不该触发下载"


class _FakePage:
    """最小页面替身：只实现 detect_store 用到的 evaluate/locator。"""

    def __init__(self, raw="", dom=None, boom=False):
        self._raw, self._dom, self._boom = raw, dom or {}, boom

    async def evaluate(self, js, *a):
        if self._boom:
            raise RuntimeError("页面已关闭")
        return self._raw

    def locator(self, sel):
        page = self

        class _Loc:
            first = property(lambda s: s)

            async def count(s):
                return 1 if sel in page._dom else 0

            async def inner_text(s):
                return page._dom.get(sel, "")

        return _Loc()


def test_detect_store_prefers_page_state():
    """优先读 rawData：店铺名挂在构建期 hash 类名上，DOM 选择器改版即失效。"""
    page = _FakePage(raw="StoreA", dom={'[class*="mallName"]': "错的名字"})

    assert asyncio.run(P.detect_store(page)) == "StoreA"


def test_detect_store_falls_back_to_dom_when_state_empty():
    """rawData 拿不到（多店账号或字段改名）时回退 DOM，取首行。"""
    page = _FakePage(raw="", dom={'[class*="mallName"]': "StoreB\n切换店铺"})

    assert asyncio.run(P.detect_store(page)) == "StoreB"


def test_detect_store_returns_empty_rather_than_guessing():
    """两条路都读不到就返回空串——service 会据此中止并要求显式指定，绝不猜落点。"""
    assert asyncio.run(P.detect_store(_FakePage(boom=True))) == ""


class _CheckboxPage:
    """页面替身：模拟 beast-core 的「隐形 input + 可见 label」结构。

    真实结构是 <input> 0×0 opacity:0（点它必然 30s 超时报 not visible），外层 label
    才可点。这里让点 input 直接抛错，钉住「必须点 label」。
    """

    def __init__(self, has_label=True, click_works=True, evaluate_boom=False):
        self.has_label, self.click_works = has_label, click_works
        self.evaluate_boom = evaluate_boom
        self.checked = False
        self.clicked: list = []
        # 记录调用顺序，用来钉「先屏蔽悬浮层再点击」
        self.calls: list = []

    async def evaluate(self, js, *a):
        if self.evaluate_boom:
            raise RuntimeError("页面已关闭")
        self.calls.append("evaluate")
        return 3

    def locator(self, sel):
        page = self
        is_label = "label" in sel

        class _Loc:
            first = property(lambda s: s)

            async def count(s):
                return 1 if (is_label and page.has_label) or not is_label else 0

            async def is_checked(s):
                return page.checked

            async def click(s):
                if not is_label:
                    raise AssertionError("点了隐形 input——它 0×0 opacity:0，实机必超时")
                page.clicked.append(sel)
                page.calls.append("click")
                if page.click_works:
                    page.checked = True

        return _Loc()

    async def wait_for_timeout(self, ms):
        return None


def test_select_all_clicks_visible_label_not_hidden_input():
    """必须点可见的 label：隐形 input 点不动，实机上就是这里卡了 30s 超时。"""
    page = _CheckboxPage()

    asyncio.run(P.select_all_on_page(page))

    assert page.checked is True
    assert "label" in page.clicked[0]


def test_select_all_verifies_state_changed():
    """点了但没勾上要当场抛：否则后面导出的是空集或漏页，静默失败最糟。"""
    page = _CheckboxPage(click_works=False)

    with pytest.raises(RuntimeError, match="状态没变成已勾选"):
        asyncio.run(P.select_all_on_page(page))


def test_select_all_is_idempotent():
    """已勾选就别再点——再点一次会取消全选。"""
    page = _CheckboxPage()
    page.checked = True

    asyncio.run(P.select_all_on_page(page))

    assert page.clicked == []


def test_select_all_disables_overlay_before_clicking():
    """先屏蔽悬浮层再点：商家助手悬浮球压在操作区上，不先屏蔽就是 30s hit-target 超时。"""
    page = _CheckboxPage()

    asyncio.run(P.select_all_on_page(page))

    assert page.calls == ["evaluate", "click"], "顺序反了等于没修"


# ---- 悬浮层遮挡（商家助手插件注入，2026-07-29 实机复现）----------------------


class _OverlayPage:
    """页面替身：只实现 evaluate，用来验 disable_pointer_overlays 的 best-effort 语义。"""

    def __init__(self, ret=0, boom=False):
        self.ret, self.boom = ret, boom
        self.args: list = []

    async def evaluate(self, js, *a):
        if self.boom:
            raise RuntimeError("Execution context was destroyed")
        self.args.append((js, a))
        return self.ret


def test_disable_pointer_overlays_targets_extension_container():
    """选择器要能命中 #temu-ass-core-ui-dashboder-root（报错里点名的那个容器）。"""
    page = _OverlayPage(ret=4)

    assert asyncio.run(P.disable_pointer_overlays(page)) == 4
    js, args = page.args[0]
    assert args == (P._POINTER_OVERLAYS,)
    assert "pointer-events" in js and "important" in js
    # 用前缀匹配而不是全等 id：`_123` 那类版本后缀会跳，见 pipeline 顶部注释
    assert "temu-ass" in P._POINTER_OVERLAYS


def test_disable_pointer_overlays_swallows_errors():
    """辅助路径失败不中断主流程：真被遮挡了后面的 click 会自己超时，报错更有诊断价值。"""
    assert asyncio.run(P.disable_pointer_overlays(_OverlayPage(boom=True))) == 0


def test_disable_pointer_overlays_sets_descendants_too():
    """必须连后代一起设：报错里拦住点击的是子节点 img，只设根节点会被子节点的 auto 盖掉。"""
    js = P._DISABLE_OVERLAY_JS

    assert "querySelectorAll('*')" in js


class _NextPage:
    """页面替身：分页容器 + 「下一页」。data-status 在点击后才翻，模拟异步换页。"""

    def __init__(self, next_cls="PGT_next_123 ", flip=True):
        self.next_cls, self.flip = next_cls, flip
        self.status = "beast-core-pagination-20-1"
        self.calls: list = []

    async def evaluate(self, js, *a):
        self.calls.append("evaluate")
        return 1

    def locator(self, sel):
        page = self
        is_next = "next" in sel

        class _Loc:
            first = property(lambda s: s)

            async def count(s):
                return 1

            async def get_attribute(s, name):
                if is_next:
                    return page.next_cls if name == "class" else None
                return page.status if name == "data-status" else None

            async def click(s):
                page.calls.append("click")
                if page.flip:
                    page.status = "beast-core-pagination-20-2"

            def locator(s, sub):
                return page.locator(sub)

        return _Loc()

    async def wait_for_selector(self, sel, timeout=None):
        return None

    async def wait_for_timeout(self, ms):
        return None


def test_goto_next_page_disables_overlay_before_clicking():
    """翻页前必须屏蔽悬浮层——实机就是在这一步卡 30s 超时后整批中止的。"""
    page = _NextPage()

    assert asyncio.run(P.goto_next_page(page)) is True
    assert page.calls == ["evaluate", "click"]


def test_goto_next_page_returns_false_on_last_page():
    """尾页（PGT_disabled）返回 False，不该点也不该抛。"""
    page = _NextPage(next_cls="PGT_next_123 PGT_disabled_123")

    assert asyncio.run(P.goto_next_page(page)) is False
    assert page.calls == []


def test_goto_next_page_raises_when_page_never_changes():
    """点了但分页状态没变要抛：静默当成翻页成功会把上一页重复抓一遍。"""
    page = _NextPage(flip=False)

    with pytest.raises(RuntimeError, match="分页状态没变"):
        asyncio.run(P.goto_next_page(page, timeout_ms=500))


def test_sweep_pages_awaits_async_on_page(monkeypatch):
    """on_page 是协程函数时必须被 await。

    service 层的进度回调是 async（要往 SSE 队列 put），早先这里直接同步调用，协程建了
    却没 await，结果 UI 上一条页进度都收不到——只在真机跑时才暴露，所以补这条守住。
    """
    import asyncio

    pages = [
        {"total": 4, "page": 1, "page_size": 2, "has_next": True},
        {"total": 4, "page": 2, "page_size": 2, "has_next": False},
    ]
    seq = iter(pages)
    cur = {"v": pages[0]}

    async def fake_read(page):
        return cur["v"]

    async def fake_next(page):
        cur["v"] = next(seq_after)

    seq_after = iter(pages[1:])
    async def fake_grab(page):
        return {f"045-{cur['v']['page']}": {"url": "u", "price": "1"}}

    async def fake_select(page):
        return None

    async def fake_count(page):
        return 2 * cur["v"]["page"]

    monkeypatch.setattr(P, "read_pagination", fake_read)
    monkeypatch.setattr(P, "grab_page_images", fake_grab)
    monkeypatch.setattr(P, "select_all_on_page", fake_select)
    monkeypatch.setattr(P, "selected_count", fake_count)
    monkeypatch.setattr(P, "goto_next_page", fake_next)

    got = []

    async def on_page(info):        # 协程回调
        got.append(info)

    res = asyncio.run(P.sweep_pages(object(), on_page=on_page))

    assert [i["page"] for i in got] == [1, 2], "两页都要推进度事件"
    assert got[-1]["selected"] == 4 and res["pages"] == 2
    assert len(res["images"]) == 2
    assert res["truncated"] is False, "正常翻完不算截断"


def _patch_sweep(monkeypatch, total=100, page_size=20):
    """给 sweep_pages 造一个「永远还有下一页」的分页，用来验 max_pages 截断。"""
    cur = {"page": 1}

    async def fake_read(page):
        return {"total": total, "page": cur["page"], "page_size": page_size,
                "has_next": cur["page"] * page_size < total}

    async def fake_next(page):
        cur["page"] += 1

    async def fake_grab(page):
        return {f"045-{cur['page']}": {"url": "u", "price": "1"}}

    # 已勾选数按「勾过几页」算，不能按当前页号——goto_next_page 会先把页号加上去，
    # 用页号会多算一页（真实语义是勾过的页数 × 每页条数）。
    picked = {"pages": 0}

    async def fake_select(page):
        picked["pages"] += 1

    async def fake_count(page):
        return page_size * picked["pages"]

    monkeypatch.setattr(P, "read_pagination", fake_read)
    monkeypatch.setattr(P, "grab_page_images", fake_grab)
    monkeypatch.setattr(P, "select_all_on_page", fake_select)
    monkeypatch.setattr(P, "selected_count", fake_count)
    monkeypatch.setattr(P, "goto_next_page", fake_next)


async def _none():
    return None


def test_sweep_pages_truncated_skips_total_check(monkeypatch):
    """max_pages 截断时不校验「已选==总数」。

    否则只要 max_pages 小于实际页数就必然抛「勾选不全」，这个冒烟参数等于死的——
    实机就是这么撞上的（248 条 5 页，--max-pages 2 直接中止）。
    """
    _patch_sweep(monkeypatch, total=100)

    res = asyncio.run(P.sweep_pages(object(), max_pages=2))

    assert res["pages"] == 2 and res["selected"] == 40
    assert res["truncated"] is True, "截断标记要传出去，别让人误读成全量"


def test_sweep_pages_still_guards_incomplete_selection(monkeypatch):
    """没被截断却已选 != 总数 → 仍要抛：少勾了就导不全，宁可整批重来。"""
    _patch_sweep(monkeypatch, total=100)

    async def short_count(page):
        return 37        # 翻完了却只勾到 37 条

    monkeypatch.setattr(P, "selected_count", short_count)

    with pytest.raises(RuntimeError, match="勾选不全"):
        asyncio.run(P.sweep_pages(object(), max_pages=5))


def test_sweep_pages_handles_zero_max_pages(monkeypatch):
    """max_pages<=0 时循环体不执行，收尾也不能 NameError。"""
    _patch_sweep(monkeypatch, total=100)

    res = asyncio.run(P.sweep_pages(object(), max_pages=0))

    assert res["pages"] == 0 and res["images"] == {}


# ---- 增量早停在 sweep 循环里的落地 -----------------------------------------


def _patch_sweep_incremental(monkeypatch, page_orders_by_page: dict, total=100,
                             page_size=20):
    """在 _patch_sweep 之上，让每页返回指定的订单号列表（模拟新→旧的列表）。"""
    cur = {"page": 1}

    async def fake_read(page):
        return {"total": total, "page": cur["page"], "page_size": page_size,
                "has_next": cur["page"] * page_size < total}

    async def fake_next(page):
        cur["page"] += 1

    async def fake_read_orders(page):
        return page_orders_by_page.get(cur["page"], [])

    _patch_sweep(monkeypatch, total=total, page_size=page_size)
    monkeypatch.setattr(P, "read_pagination", fake_read)
    monkeypatch.setattr(P, "goto_next_page", fake_next)
    monkeypatch.setattr(P, "read_page_orders", fake_read_orders)
    return cur


def test_sweep_stops_early_when_caught_up(monkeypatch):
    """第 1 页有新单、第 2 页全已登记 → 停在第 2 页，不再翻剩下 3 页。"""
    _patch_sweep_incremental(monkeypatch, {
        1: _page("PO-new1", "PO-new2"),
        2: _page("PO-old1", "PO-old2"),
        3: _page("PO-old3"),
    })

    res = asyncio.run(P.sweep_pages(
        object(), max_pages=5, known_order_nos={"PO-old1", "PO-old2", "PO-old3"}
    ))

    assert res["pages"] == 2, "追上就停，不翻到底"
    assert res["stopped_early"] is True
    assert res["truncated"] is False, "早停不是截断，语义相反，绝不能混"
    assert res["fell_back"] is False


def test_sweep_early_stop_skips_total_check(monkeypatch):
    """早停时已选 < 总数是必然的，不能抛「勾选不全」。"""
    _patch_sweep_incremental(monkeypatch, {
        1: _page("PO-new1"),
        2: _page("PO-old1"),
    }, total=100)

    res = asyncio.run(P.sweep_pages(
        object(), max_pages=5, known_order_nos={"PO-old1"}
    ))

    assert res["stopped_early"] is True and res["selected"] < res["total"]


def test_sweep_stops_on_sparse_watermark(monkeypatch):
    """稀疏水位回归：水位只有 1 条、它在第 4 页中间 → 停在第 4 页，不翻到底。

    2026-08-04 实机 bug 的最小复现（原先翻了 63 页）：旧判据要整页全已登记才停，
    水位 1 条时永不成立；且第 4 页「新单→已登记→新单」的排布会被交错检测判成排序异常
    而永久退全量。两个缺陷叠加使增量在新建表上完全失效。
    """
    _patch_sweep_incremental(monkeypatch, {
        1: _page("PO-n1", "PO-n2"),
        2: _page("PO-n3", "PO-n4"),
        3: _page("PO-n5", "PO-n6"),
        4: _page("PO-n7", "PO-old1", "PO-n8"),   # 已登记单夹在新单中间
        5: _page("PO-n9"),
    })

    res = asyncio.run(P.sweep_pages(
        object(), max_pages=5, known_order_nos={"PO-old1"},
    ))

    assert res["pages"] == 4, "遇到已登记单就停，不该翻到第 5 页"
    assert res["stopped_early"] is True
    assert res["fell_back"] is False, "稀疏水位是正常状态，不该报退全量"


def test_sweep_unreadable_page_does_not_poison_batch(monkeypatch):
    """某页读不到订单号 → 只跳过本页，后续页照常判早停，绝不整批退全量。

    read_page_orders 是 best-effort（DOM 瞬时读失败返回空表），它的语义是「本页不参与
    判断」。改动前调用点把这种情形也置了 desc_broken，一次瞬时失败就让整批增量失效。
    """
    _patch_sweep_incremental(monkeypatch, {
        1: _page("PO-n1"),
        2: [],                    # 读不到
        3: _page("PO-old1"),
        4: _page("PO-old2"),
        5: _page("PO-old3"),
    })

    res = asyncio.run(P.sweep_pages(
        object(), max_pages=5, known_order_nos={"PO-old1", "PO-old2", "PO-old3"},
    ))

    assert res["pages"] == 3, "第 2 页读不到只跳过本页，第 3 页照样早停"
    assert res["stopped_early"] is True
    assert res["fell_back"] is False


def test_sweep_without_watermark_keeps_full_behavior(monkeypatch):
    """不传水位＝改动前的行为：翻到底、不早停、不标 fell_back。"""
    _patch_sweep_incremental(monkeypatch, {1: _page("PO-1")}, total=40)

    res = asyncio.run(P.sweep_pages(object(), max_pages=5))

    assert res["stopped_early"] is False and res["fell_back"] is False
    assert res["pages"] == 2 and res["truncated"] is False


def test_sweep_reports_new_and_known_counts_per_page(monkeypatch):
    """页进度要带本页新/已登记计数，UI 上才看得出增量在起作用。"""
    _patch_sweep_incremental(monkeypatch, {
        1: _page("PO-new1", "PO-new2"),
        2: _page("PO-new3", "PO-old1"),
    })
    got = []

    async def on_page(info):
        got.append(info)

    asyncio.run(P.sweep_pages(
        object(), on_page=on_page, max_pages=5,
        known_order_nos={"PO-old1", "PO-old2"},
    ))

    assert (got[0]["new_on_page"], got[0]["known_on_page"]) == (2, 0)
    assert (got[1]["new_on_page"], got[1]["known_on_page"]) == (1, 1)
